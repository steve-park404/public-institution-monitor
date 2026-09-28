import csv, json, os, sys
from datetime import datetime
from pathlib import Path
from openpyxl import load_workbook

CSV_FILE = "기관별_상태.csv"
DIAG_FILE = "diagnostics.json"
TARGET_FILES = ["monitor_targets.xlsx", "url.xlsx", "targets.xlsx"]

FIELDS = [
    "기관명","기관유형","홈페이지URL","게시판URL","상태","상태설명",
    "게시물후보","상세확인","기존확인건너뜀","최근게시물","제목매칭","본문매칭",
    "상세오류","게시판후보수","게시판조회수","게시판검증수","선정점수",
    "재시도대상","최종확인시각","홈페이지복구시도","홈페이지복구성공","복구URL",
    "게시판판정","게시판판정근거","최고후보URL","최고후보점수"
]

def norm(x):
    return "" if x is None else str(x).strip()

def load_targets():
    target = next((p for p in TARGET_FILES if Path(p).exists()), None)
    if not target:
        raise FileNotFoundError(f"기관 목록 파일이 없습니다: {TARGET_FILES}")
    wb = load_workbook(target, read_only=True, data_only=True)
    ws = wb["355기관"] if "355기관" in wb.sheetnames else wb[wb.sheetnames[0]]
    rows = list(ws.iter_rows(values_only=True))
    headers = [norm(x) for x in rows[0]]
    def col(*names):
        for n in names:
            if n in headers:
                return headers.index(n)
        return None
    ni, hi, bi, ti = col("기관명","기관"), col("홈페이지URL","URL","홈페이지"), col("공지게시판URL","공지사항URL","게시판URL"), col("기관유형","기관 유형","유형","기관구분","구분")
    if ni is None or hi is None:
        raise ValueError(f"기관명/홈페이지URL 열 없음: {headers}")
    out=[]
    for row in rows[1:]:
        if not row: continue
        name=norm(row[ni]) if ni<len(row) else ""
        home=norm(row[hi]) if hi<len(row) else ""
        board=norm(row[bi]) if bi is not None and bi<len(row) else ""
        typ=norm(row[ti]) if ti is not None and ti<len(row) else ""
        if name and home:
            out.append({"기관명":name,"기관유형":typ,"홈페이지URL":home,"게시판URL":board})
    return out

def status_desc(status):
    return {
        "CACHED_OR_SEED":"기존 캐시/지정 게시판 확인",
        "DISCOVERED":"홈페이지에서 게시판을 새로 발견",
        "HOME_ERROR":"홈페이지 접속 실패",
        "NO_CANDIDATE":"게시판 후보를 찾지 못함",
        "CANDIDATE_NOT_VERIFIED":"게시판 후보는 있으나 실제 게시판 검증 실패",
        "PROCESS_ERROR":"기관 처리 중 오류",
        "BOARD_FETCH_ERROR":"게시판 접속 실패",
        "UNPROCESSED_TIMEOUT":"실행시간 초과로 미처리",
        "UNPROCESSED":"미처리"
    }.get(status, status or "미상")

def num(x):
    return x if isinstance(x,(int,float)) else (int(x) if str(x).isdigit() else 0)

def main():
    if not Path(DIAG_FILE).exists():
        raise FileNotFoundError("diagnostics.json이 없습니다.")
    diag=json.loads(Path(DIAG_FILE).read_text(encoding="utf-8"))
    targets=load_targets()
    by_name={x.get("기관명"):x for x in diag.get("board_details",[]) if x.get("기관명")}
    now=diag.get("updated_at") or datetime.now().astimezone().isoformat()

    rows=[]
    for t in targets:
        name=t["기관명"]
        r=by_name.get(name)
        if r is None:
            status="UNPROCESSED_TIMEOUT" if diag.get("timed_out") else "UNPROCESSED"
            rows.append({
                "기관명":name,"기관유형":t["기관유형"],"홈페이지URL":t["홈페이지URL"],
                "게시판URL":t["게시판URL"],"상태":status,"상태설명":status_desc(status),
                "재시도대상":"예","최종확인시각":now
            })
            continue
        status=r.get("status") or r.get("diag_status") or "UNKNOWN"
        rows.append({
            "기관명":name,
            "기관유형":r.get("기관유형",t["기관유형"]),
            "홈페이지URL":t["홈페이지URL"],
            "게시판URL":r.get("board") or t["게시판URL"],
            "상태":status,
            "상태설명":status_desc(status),
            "게시물후보":num(r.get("post_candidates")),
            "상세확인":num(r.get("detail_checked")),
            "기존확인건너뜀":num(r.get("skipped_previously_checked")),
            "최근게시물":num(r.get("recent_posts")),
            "제목매칭":num(r.get("title_matches")),
            "본문매칭":num(r.get("body_matches")),
            "상세오류":num(r.get("detail_errors")),
            "게시판후보수":num(r.get("candidate_count")),
            "게시판조회수":num(r.get("fetched_count")),
            "게시판검증수":num(r.get("verified_count")),
            "선정점수":r.get("selected_score",""),
            "재시도대상":"예" if status in {"HOME_ERROR","NO_CANDIDATE","CANDIDATE_NOT_VERIFIED","BOARD_FETCH_ERROR","PROCESS_ERROR"} else "아니오",
            "최종확인시각":now,
            "게시판판정":r.get("diag_status",""),
            "게시판판정근거":"",
            "최고후보URL":"",
            "최고후보점수":""
        })
    # Exact target count is a hard invariant.
    if len(rows)!=len(targets) or len(rows)!=355:
        raise RuntimeError(f"CSV 행 수 오류: {len(rows)} / 기대 355")
    tmp=Path(CSV_FILE+".tmp")
    with tmp.open("w",encoding="utf-8-sig",newline="") as f:
        w=csv.DictWriter(f,fieldnames=FIELDS,extrasaction="ignore")
        w.writeheader()
        for row in rows:
            w.writerow({k:row.get(k,"") for k in FIELDS})
    os.replace(tmp,CSV_FILE)
    size=Path(CSV_FILE).stat().st_size
    print(f"[CSV] {CSV_FILE} 생성 완료: {len(rows)}개 행, {size:,} bytes")

if __name__=="__main__":
    main()
