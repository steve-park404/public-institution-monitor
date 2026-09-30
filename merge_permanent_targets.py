from pathlib import Path
import openpyxl
BASE=["monitor_targets.xlsx","url.xlsx","targets.xlsx"]
EXTRA="유사기관_추가모니터링_후보목록_URL입력완료 (2).xlsx"
OUT="monitor_targets_통합.xlsx"
def n(v): return str(v).strip() if v is not None else ""
base=next((Path(x) for x in BASE if Path(x).exists()),None)
if not base: raise SystemExit("기존 355기관 Excel을 찾지 못했습니다.")
if not Path(EXTRA).exists(): raise SystemExit("유사기관 Excel을 찾지 못했습니다.")
wb=openpyxl.load_workbook(base)
ws=wb["355기관"] if "355기관" in wb.sheetnames else wb[wb.sheetnames[0]]
headers=[n(c.value) for c in ws[1]]
h={v:i+1 for i,v in enumerate(headers) if v}
for col in ["기관명","홈페이지URL","공지게시판URL","기관유형","모니터링출처"]:
    if col not in h:
        h[col]=ws.max_column+1; ws.cell(1,h[col],col)
existing=set()
for row in ws.iter_rows(min_row=2,values_only=True):
    existing.add((n(row[h["기관명"]-1]),n(row[h["홈페이지URL"]-1]).rstrip("/")))
ewb=openpyxl.load_workbook(EXTRA,read_only=True,data_only=True)
ews=ewb["추가후보"]; it=ews.iter_rows(values_only=True); ehdr=next(it)
eh={n(v):i for i,v in enumerate(ehdr)}
added=dup=no_board=0
for row in it:
    name=n(row[eh["기관명"]]); home=n(row[eh["홈페이지URL"]]); board=n(row[eh["게시판URL"]])
    org=n(row[eh["기관유형"]]) if "기관유형" in eh else ""
    if not name or not home: continue
    if not board: no_board+=1; continue
    key=(name,home.rstrip("/"))
    if key in existing: dup+=1; continue
    r=ws.max_row+1
    ws.cell(r,h["기관명"],name); ws.cell(r,h["홈페이지URL"],home)
    ws.cell(r,h["공지게시판URL"],board); ws.cell(r,h["기관유형"],org)
    ws.cell(r,h["모니터링출처"],"수동확인 유사기관")
    existing.add(key); added+=1
wb.save(OUT)
total=ws.max_row-1
print(f"최종 통합 대상: {total}개 / 추가: {added} / 중복: {dup} / 게시판없음: {no_board}")
if total<500: raise SystemExit("통합 대상이 500개 미만입니다.")
