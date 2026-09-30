# -*- coding: utf-8 -*-
"""
V8.23 유사기관 URL 병합
- 기존 monitor.py의 검색 로직은 수정하지 않음
- 기존 355기관 원본 Excel에 사용자가 수동 확인한 유사기관 게시판 URL을 추가
- 기관명 중복은 1개로 유지
"""

from pathlib import Path
import re
import shutil
from datetime import datetime
import openpyxl

BASE_CANDIDATES = ["monitor_targets.xlsx", "url.xlsx", "targets.xlsx"]
EXTRA_CANDIDATES = [
    "유사기관_추가모니터링_후보목록_URL입력완료 (2).xlsx",
    "유사기관_추가모니터링_후보목록_URL입력완료.xlsx",
    "유사기관_추가모니터링_후보목록.xlsx",
]

def norm(x):
    return re.sub(r"\s+", "", str(x or "")).strip()

def valid_url(x):
    return bool(re.match(r"^https?://", str(x or "").strip(), re.I))

def find_existing(names):
    for n in names:
        p = Path(n)
        if p.exists():
            return p
    return None

def headers(ws):
    return {norm(c.value): i for i, c in enumerate(ws[1], start=1)}

def main():
    base = find_existing(BASE_CANDIDATES)
    extra = find_existing(EXTRA_CANDIDATES)

    if not base:
        raise SystemExit(
            "기존 기관목록 Excel을 찾지 못했습니다. "
            f"다음 중 하나가 필요합니다: {BASE_CANDIDATES}"
        )
    if not extra:
        raise SystemExit("유사기관 추가 Excel을 찾지 못했습니다.")

    # 이미 병합된 경우에도 안전하게 재실행 가능
    wb = openpyxl.load_workbook(base)
    ws = wb["355기관"] if "355기관" in wb.sheetnames else wb[wb.sheetnames[0]]
    hm = headers(ws)

    if "기관명" not in hm or "홈페이지URL" not in hm:
        raise SystemExit(f"기존 Excel의 기관명/홈페이지URL 열을 찾지 못했습니다: {list(hm)}")

    # 필요한 열이 없으면 추가
    for colname in ["공지게시판URL", "기관유형"]:
        if colname not in hm:
            idx = ws.max_column + 1
            ws.cell(1, idx).value = colname
            hm[colname] = idx

    # 기존 기관명 색인
    existing = {}
    for r in range(2, ws.max_row + 1):
        name = norm(ws.cell(r, hm["기관명"]).value)
        if name:
            existing[name] = r

    ewb = openpyxl.load_workbook(extra, read_only=True, data_only=True)
    ews = ewb["추가후보"] if "추가후보" in ewb.sheetnames else ewb[ewb.sheetnames[0]]
    eh = headers(ews)

    if "기관명" not in eh or "게시판URL" not in eh:
        raise SystemExit(f"추가 Excel의 기관명/게시판URL 열을 찾지 못했습니다: {list(eh)}")

    added = 0
    supplemented = 0
    skipped = 0

    for r in range(2, ews.max_row + 1):
        name = norm(ews.cell(r, eh["기관명"]).value)
        board = str(ews.cell(r, eh["게시판URL"]).value or "").strip()
        home = str(ews.cell(r, eh["홈페이지URL"]).value or "").strip() if "홈페이지URL" in eh else ""
        org_type = str(ews.cell(r, eh["기관유형"]).value or "").strip() if "기관유형" in eh else ""

        # 게시판 URL을 실제로 수동 입력한 기관만 편입
        if not name or not valid_url(board):
            continue

        if name in existing:
            rr = existing[name]
            old_board = str(ws.cell(rr, hm["공지게시판URL"]).value or "").strip()
            changed = False

            if not valid_url(old_board):
                ws.cell(rr, hm["공지게시판URL"]).value = board
                changed = True

            if "홈페이지URL" in hm and not str(ws.cell(rr, hm["홈페이지URL"]).value or "").strip() and home:
                ws.cell(rr, hm["홈페이지URL"]).value = home
                changed = True

            if changed:
                supplemented += 1
            else:
                skipped += 1
            continue

        rr = ws.max_row + 1
        ws.cell(rr, hm["기관명"]).value = name
        ws.cell(rr, hm["홈페이지URL"]).value = home or board
        ws.cell(rr, hm["공지게시판URL"]).value = board
        ws.cell(rr, hm["기관유형"]).value = org_type
        existing[name] = rr
        added += 1

    wb.save(base)

    print("=== V8.23 유사기관 URL 병합 ===")
    print(f"기존 대상 파일: {base}")
    print(f"신규 기관 추가: {added}")
    print(f"기존 기관 URL 보완: {supplemented}")
    print(f"기존값 유지: {skipped}")
    print(f"최종 기관 수: {ws.max_row - 1}")

if __name__ == "__main__":
    main()
